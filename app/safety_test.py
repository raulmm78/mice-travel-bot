from __future__ import annotations

import os
import tempfile
from pathlib import Path
from datetime import date
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from process_emails import (
    TravelRequest,
    is_travel_request_email,
    write_basic_excel,
    write_nn_excel,
    workbook_digest,
    save_workbook_safely,
)
import process_emails


def assert_true(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def request(nombre: str, dni: str, email_id: str = "mail") -> TravelRequest:
    return TravelRequest(
        email_id=email_id,
        source_message_id=f"<{email_id}@example.com>",
        evento="Congreso IMS",
        nombre=nombre,
        apellidos="TEST",
        dni=dni,
        email=f"{nombre.lower()}@example.com",
        origen="Madrid",
        destino="Rio",
        fecha_viaje="2026-09-28",
        estado="ok",
    )


def test_filter() -> None:
    assert_true(
        is_travel_request_email(
            "RV: NUEVA SOLICITUD: Congreso IMS - TEST",
            "Texto previo\nDATOS PARA PETICIONES DE INVITADOS\nNOMBRE: TEST",
        ),
        "No reconoce reenvios RV/FW validos con formulario",
    )
    assert_true(
        not is_travel_request_email("NUEVA SOLICITUD: otra cosa", "Correo sin formulario"),
        "No esta filtrando correos sin formulario",
    )


def test_basic_append() -> None:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "basic.xlsx"
        old = request("ANTIGUO", "11111111H", "old")
        new = request("NUEVO", "22222222J", "new")
        write_basic_excel([old], path, "TEST")
        write_basic_excel([old, new], path, "TEST")
        write_basic_excel([old, new], path, "TEST")
        ws = load_workbook(path)["Totales"]
        dni_values = [ws.cell(row=row, column=5).value for row in range(3, ws.max_row + 1)]
        assert_true(dni_values.count("11111111H") == 1, "La fila antigua se ha duplicado o perdido")
        assert_true(dni_values.count("22222222J") == 1, "La fila nueva no se ha anadido una sola vez")


def test_nn_append() -> None:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "existing.xlsx"
        template_path = Path(temp) / "template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Totales"
        for col in range(1, 46):
            sheet.cell(row=3, column=col).value = f"H{col}"
        sheet["D4"] = "CONGRESO IMS"
        sheet["M4"] = "ANTIGUO"
        sheet["N4"] = "TEST"
        sheet["O4"] = "11111111H"
        sheet["U4"] = "antiguo@example.com"
        sheet["X4"] = date(2026, 9, 28)
        sheet["AQ5"] = "NOTA MANUAL"
        workbook.save(path)
        workbook.save(template_path)

        os.environ["NN_TEMPLATE_PATH"] = str(template_path)
        old = request("ANTIGUO", "11111111H", "old")
        new = request("NUEVO", "22222222J", "new")
        write_nn_excel([old, new], path, "TEST")
        write_nn_excel([old, new], path, "TEST")

        sheet = load_workbook(path)["Totales"]
        assert_true(sheet["M4"].value == "ANTIGUO", "La fila antigua ha cambiado")
        assert_true(sheet["O4"].value == "11111111H", "El DNI antiguo ha cambiado")
        assert_true(sheet["AQ5"].value == "NOTA MANUAL", "Se ha borrado una celda manual")
        assert_true(sheet["M6"].value == "NUEVO", "No se ha anadido la fila nueva debajo")
        assert_true(sheet["O6"].value == "22222222J", "El DNI nuevo no esta en la fila esperada")
        assert_true(sheet["O7"].value in (None, ""), "Se ha duplicado la fila nueva")
        backups = list((path.parent / "_bot_backups").glob("*.xlsx"))
        assert_true(len(backups) == 1, "Falta la copia del Excel anterior")
        assert_true(load_workbook(backups[0])["Totales"]["AQ5"].value == "NOTA MANUAL",
                    "La copia de seguridad no conserva la nota manual")


def test_conflict_stops_before_replace() -> None:
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "original.xlsx"
        old_book = Workbook()
        old_book.active["A1"] = "ORIGINAL"
        old_book.save(path)
        original_digest = workbook_digest(path)
        new_book = Workbook()
        new_book.active["A1"] = "CAMBIO"
        try:
            save_workbook_safely(new_book, path, "digest-desactualizado")
        except RuntimeError:
            pass
        else:
            raise AssertionError("No se detuvo el guardado ante un cambio concurrente")
        assert_true(workbook_digest(path) == original_digest, "El conflicto modifico el Excel")
        try:
            save_workbook_safely(new_book, path)
        except FileExistsError:
            pass
        else:
            raise AssertionError("Crear Excel sustituyo un archivo ya existente")
        assert_true(workbook_digest(path) == original_digest, "Crear Excel modifico un archivo existente")


def test_global_and_event_flow() -> None:
    with tempfile.TemporaryDirectory() as temp:
        root = Path(temp)
        template = root / "plantilla.xlsx"
        event_template = root / "plantilla_eventos.xlsx"
        global_path = root / "global.xlsx"
        events = root / "eventos"
        events.mkdir()
        workbook = Workbook()
        workbook.active.title = "Totales"
        workbook.active["M4"] = "PERSONA DE PLANTILLA"
        workbook.active["O4"] = "99999999R"
        workbook.save(template)
        template_digest = workbook_digest(template)
        event_book = Workbook()
        event_sheet = event_book.active
        event_sheet.title = "Listado"
        event_book.create_sheet("Condiciones")
        for column, label in {
            "M": "NOMBRE", "N": "APELLIDOS", "O": "DNI", "X": "FECHA INICIO SERVICIO",
            "AN": "ALERGIAS", "AQ": "OBSERVACIONES",
        }.items():
            event_sheet[f"{column}3"] = label
        event_sheet["A4"] = 1
        for row_num in range(4, 8):
            event_sheet[f"BE{row_num}"] = f"=SUM(AR{row_num}:BD{row_num})"
        event_sheet["BI5"] = 0
        event_sheet["DM4"] = "AIR ALGERIE"
        event_sheet["M4"].border = Border(left=Side(style="thin"))
        event_sheet["M4"].fill = PatternFill("solid", fgColor="FFFFFF")
        validation = DataValidation(type="list", formula1='"A,B"')
        event_sheet.add_data_validation(validation)
        validation.add("M4:M7")
        event_sheet.conditional_formatting.add("M4:M7", CellIsRule(operator="equal", formula=["0"]))
        event_book.save(event_template)
        old = request("ANTIGUO", "11111111H", "old")
        new = request("NUEVO", "22222222J", "new")
        other = request("OTRO", "33333333P", "other")
        other.evento = "Congreso EULAR"
        other.conex_ida = "DIRECTO"
        other.hotel_in = "2026-09-28"
        other.restricciones_alimentarias = "SIN LACTOSA"
        existing_event = events / "NO ENVIAR -----LISTADO CONGRESO IMS 28 SEP.xlsx"
        env = {
            "NN_TEMPLATE_PATH": str(template),
            "EVENT_TEMPLATE_PATH": str(event_template),
            "OUTPUT_DIR": str(root),
            "XLSX_PATH": str(global_path),
            "EVENT_OUTPUT_DIR": str(events),
            "EVENT_ROUTES_PATH": str(root / "routes.json"),
        }
        with patch.dict(os.environ, env):
            try:
                write_nn_excel([request("ERROR", "88888888Q")], template, "NO DEBE ESCRIBIR")
            except ValueError:
                pass
            else:
                raise AssertionError("Se permitio escribir sobre la plantilla")
            write_nn_excel([old], global_path, "GLOBAL")
            write_nn_excel([old], existing_event, "Congreso IMS")
            process_emails.assign_event_excel("Congreso IMS", existing_event)
            process_emails.write_outputs([old, new, other])
            assert_true(workbook_digest(template) == template_digest, "La plantilla ha cambiado")
            assert_true(existing_event.exists(), "Se ha eliminado el Excel de evento")
            suggested = process_emails.suggested_event_workbook_path("Congreso EULAR", [other])
            assert_true(not suggested.exists(), "Se ha creado un Excel antes de confirmar el evento")
            with patch.object(process_emails, "process_all", return_value=[old, new, other]):
                created = process_emails.create_event_excel("Congreso EULAR")
            assert_true(created == suggested and created.exists(), "No se ha creado el Excel sugerido")
            process_emails.write_outputs([old, new, other])
            next_attendee = request("CUARTO", "44444444L", "fourth")
            next_attendee.evento = "Congreso EULAR"
            process_emails.write_nn_excel([other, next_attendee], suggested, "Congreso EULAR",
                                          process_emails.event_template_path())

        global_sheet = load_workbook(global_path)["Totales"]
        event_sheet = load_workbook(existing_event)["Totales"]
        new_book = load_workbook(suggested)
        new_sheet = new_book["Listado"]
        source_sheet = load_workbook(event_template)["Listado"]
        assert_true([global_sheet[f"O{row}"].value for row in (4, 5, 6)] ==
                    ["11111111H", "22222222J", "33333333P"], "El global perdio o duplico filas")
        assert_true([event_sheet[f"O{row}"].value for row in (4, 5)] ==
                    ["11111111H", "22222222J"], "El Excel de evento perdio o duplico filas")
        assert_true(new_sheet["O4"].value == "33333333P" and new_sheet["O5"].value == "44444444L"
                    and new_sheet["O6"].value is None, "El evento nuevo no anade debajo sin duplicar")
        assert_true(new_sheet["Y4"].value == "DIRECTO", "La conexion de ida fue a otra columna")
        assert_true(new_sheet["AD4"].value is not None, "La fecha de entrada del hotel no esta en IN")
        assert_true(new_sheet["AN4"].value == "SIN LACTOSA", "La alergia no esta en su columna")
        assert_true(new_sheet["BE4"].value == source_sheet["BE4"].value,
                    "Se ha perdido una formula de la plantilla")
        assert_true(new_sheet["DM4"].value == source_sheet["DM4"].value,
                    "Se ha perdido la lista auxiliar de la plantilla")
        assert_true(new_sheet["M4"].fill.patternType == source_sheet["M4"].fill.patternType
                    and new_sheet["M4"].border.left.style == source_sheet["M4"].border.left.style,
                    "Se ha cambiado el formato base de la plantilla")
        assert_true("Condiciones" in new_book.sheetnames, "Falta una hoja de la plantilla")
        assert_true(len(new_sheet.data_validations.dataValidation) == len(source_sheet.data_validations.dataValidation),
                    "Se han perdido validaciones de la plantilla")
        assert_true(len(new_sheet.conditional_formatting) == len(source_sheet.conditional_formatting),
                    "Se han perdido formatos condicionales de la plantilla")


def test_supplied_template_when_available() -> None:
    template = process_emails.BUNDLED_NN_TEMPLATE_PATH
    if not template.exists():
        return
    with tempfile.TemporaryDirectory() as temp:
        path = Path(temp) / "nuevo_evento.xlsx"
        attendee = request("CLIENTE", "55555555M", "supplied")
        attendee.conex_ida = "DIRECTO"
        attendee.restricciones_alimentarias = "SIN GLUTEN"
        write_nn_excel([attendee], path, "Congreso IMS", template)
        source = load_workbook(template)["Listado"]
        output = load_workbook(path)
        sheet = output["Listado"]
        assert_true(sheet["O4"].value == "55555555M", "La plantilla real no usa la primera fila")
        assert_true(sheet["Y4"].value == "DIRECTO" and sheet["AN4"].value == "SIN GLUTEN",
                    "La plantilla real tiene columnas desalineadas")
        assert_true(sheet["BE4"].value == source["BE4"].value, "La plantilla real perdio formulas")
        assert_true(sheet["DM4"].value == source["DM4"].value, "La plantilla real perdio listas")
        assert_true(len(sheet.conditional_formatting) == len(source.conditional_formatting),
                    "La plantilla real perdio formato condicional")
        assert_true("Condiciones" in output.sheetnames, "La plantilla real perdio hojas")


def main() -> None:
    test_filter()
    test_basic_append()
    test_nn_append()
    test_conflict_stops_before_replace()
    test_global_and_event_flow()
    test_supplied_template_when_available()
    print("OK - Seguridad Excel verificada")
    print("- Filtra correos que no son formulario")
    print("- Acepta reenvios RV/FW si contienen el formulario")
    print("- No borra filas antiguas")
    print("- No duplica la misma solicitud al repetir proceso")


if __name__ == "__main__":
    main()
